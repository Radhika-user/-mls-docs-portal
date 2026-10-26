import os
import io
import time
import shutil
import base64
from datetime import datetime
import pandas as pd
import streamlit as st

# Optional parsers
try:
    import docx
    from docx.oxml.ns import qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

# ---------------------------------------------------------
# Page Configuration & Styling
# ---------------------------------------------------------
st.set_page_config(
    page_title="MLS & Research Docs Portal",
    page_icon="📚",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Styling configuration
app_bg_css = """
[data-testid="stAppViewContainer"] {
    background-color: #F1F5F9 !important;
}
[data-testid="stSidebar"] {
    background-color: #FFFFFF !important;
    border-right: 1.5px solid #CBD5E1 !important;
}
"""

st.markdown(f"""
<style>
    {app_bg_css}
    .portal-hero {{
        background: linear-gradient(135deg, #0F172A 0%, #1E3A8A 60%, #1D4ED8 100%);
        border-radius: 12px;
        padding: 24px 30px;
        margin-bottom: 22px;
        box-shadow: 0 6px 20px rgba(30, 58, 138, 0.18);
        border: 1px solid rgba(255, 255, 255, 0.1);
    }}
    .main-title {{
        font-size: 1.85rem;
        font-weight: 700;
        color: #1E3A8A;
        margin-bottom: 0.2rem;
    }}
    .sub-title {{
        color: #475569;
        font-size: 0.95rem;
        margin-bottom: 1.2rem;
    }}
    .metric-card {{
        background-color: #FFFFFF;
        border-radius: 10px;
        padding: 16px 14px;
        border: 1px solid #CBD5E1;
        border-top: 3px solid #1E3A8A;
        text-align: center;
        box-shadow: 0 4px 10px rgba(30, 58, 138, 0.05);
    }}
    .doc-meta {{
        background: #FFFFFF;
        padding: 16px 20px;
        border-radius: 10px;
        margin-bottom: 18px;
        border: 1px solid #CBD5E1;
        border-left: 5px solid #1E3A8A;
        box-shadow: 0 3px 8px rgba(30, 58, 138, 0.04);
    }}
    .doc-reader-canvas {{
        background: #FFFFFF;
        border: 1px solid #CBD5E1;
        border-radius: 12px;
        padding: 30px 35px;
        box-shadow: 0 4px 16px rgba(0,0,0,0.04);
        margin-top: 15px;
    }}
    [data-testid="stForm"] {{
        background: #FFFFFF !important;
        border-radius: 14px !important;
        padding: 32px 34px 26px 34px !important;
        border: 1px solid #CBD5E1 !important;
        box-shadow: 0 20px 45px rgba(15, 23, 42, 0.18) !important;
    }}
    [data-testid="stForm"] label {{
        font-weight: 600 !important;
        color: #0F172A !important;
        font-size: 0.92rem !important;
    }}
    [data-testid="stForm"] input {{
        background-color: #F8FAFC !important;
        border: 1.5px solid #CBD5E1 !important;
        border-radius: 8px !important;
        color: #0F172A !important;
    }}
    [data-testid="stForm"] input:focus {{
        border-color: #1E3A8A !important;
        background-color: #FFFFFF !important;
        box-shadow: 0 0 0 2px rgba(30, 58, 138, 0.15) !important;
    }}
    .stExpander {{
        background: #FFFFFF !important;
        border: 1px solid #CBD5E1 !important;
        border-left: 4px solid #1E3A8A !important;
        border-radius: 8px !important;
        box-shadow: 0 2px 5px rgba(30, 58, 138, 0.03) !important;
        margin-bottom: 9px !important;
    }}
    button[kind="primary"], [data-testid="stBaseButton-primary"], [data-testid="stFormSubmitButton"] button {{
        background-color: #1D4ED8 !important;
        border: 1px solid #1D4ED8 !important;
        color: #FFFFFF !important;
        font-weight: 600 !important;
        box-shadow: 0 4px 12px rgba(29, 78, 216, 0.28) !important;
    }}
    button[kind="primary"]:hover, [data-testid="stBaseButton-primary"]:hover, [data-testid="stFormSubmitButton"] button:hover {{
        background-color: #1E40AF !important;
        border-color: #1E40AF !important;
        color: #FFFFFF !important;
    }}
    button[kind="secondary"], [data-testid="stBaseButton-secondary"] {{
        border: 1px solid #CBD5E1 !important;
        color: #1E293B !important;
        font-weight: 500 !important;
        background-color: #FFFFFF !important;
    }}
    button[kind="secondary"]:hover, [data-testid="stBaseButton-secondary"]:hover {{
        border-color: #1D4ED8 !important;
        color: #1D4ED8 !important;
    }}
    .badge {{
        display: inline-block;
        padding: 3px 8px;
        border-radius: 6px;
        font-size: 0.72rem;
        font-weight: 700;
        letter-spacing: 0.02em;
    }}
    .badge-docx {{ background-color: #EFF6FF; color: #1D4ED8; border: 1px solid #BFDBFE; }}
    .badge-xlsx {{ background-color: #F0FDF4; color: #15803D; border: 1px solid #BBF7D0; }}
    .badge-xls  {{ background-color: #FEFCE8; color: #A16207; border: 1px solid #FEF08A; }}
    .badge-pptx {{ background-color: #FEF2F2; color: #B91C1C; border: 1px solid #FECACA; }}
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# Constants & Scanning
# ---------------------------------------------------------
DEFAULT_DOCS_ROOT = r"\\pfbafileserver\File Server\Admin Ops\Research\MLS\Procedures"

MIME_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".pdf": "application/pdf",
    ".txt": "text/plain"
}

def get_mime_type(ext: str) -> str:
    return MIME_TYPES.get(ext.lower(), "application/octet-stream")

@st.cache_data(ttl=30)
def scan_documents_folder(base_path: str):
    if not os.path.exists(base_path):
        return []
    
    file_records = []
    for root, dirs, files in os.walk(base_path):
        for file in files:
            # Ignore Microsoft Office temporary lock files and hidden system files
            if file.startswith("~$") or file.startswith("~") or file.startswith("."):
                continue
            
            ext = os.path.splitext(file)[1].lower()
            if ext in ['.docx', '.xlsx', '.xls', '.pptx', '.pdf', '.txt']:
                full_path = os.path.join(root, file)
                
                try:
                    stat = os.stat(full_path)
                    # Exclude empty or stub temporary files (< 512 bytes)
                    if stat.st_size < 512:
                        continue
                    size_kb = stat.st_size / 1024
                    mod_time = datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M')
                except Exception:
                    continue
                
                rel_path = os.path.relpath(full_path, base_path)
                parts = rel_path.split(os.sep)
                if len(parts) > 1:
                    category = parts[0]
                    sub_category = os.sep.join(parts[:-1])
                else:
                    category = "Root Procedures & SOPs"
                    sub_category = "General"
                
                file_records.append({
                    "name": file,
                    "ext": ext,
                    "full_path": full_path,
                    "rel_path": rel_path,
                    "category": category,
                    "sub_category": sub_category,
                    "size_kb": size_kb,
                    "size_mb": size_kb / 1024,
                    "mod_time": mod_time
                })
    
    file_records.sort(key=lambda x: (x["category"], x["name"]))
    return file_records

# ---------------------------------------------------------
# Parsers (Cached)
# ---------------------------------------------------------
def extract_element_images(elem, doc):
    """Safely extracts embedded screenshots/diagrams (DrawingML and VML) without crashing."""
    imgs = []
    try:
        blips = elem.xpath('.//*[local-name()="blip"]/@*[local-name()="embed"]')
        vmls = elem.xpath('.//*[local-name()="imagedata"]/@*[local-name()="id"]')
        for rId in blips + vmls:
            try:
                if rId in doc.part.related_parts:
                    part = doc.part.related_parts[rId]
                    if hasattr(part, 'blob'):
                        imgs.append(part.blob)
            except Exception:
                pass
    except Exception:
        pass
    return imgs

@st.cache_data(show_spinner=False)
def parse_docx(file_path: str):
    if not DOCX_AVAILABLE:
        return {"error": "python-docx library is not installed."}
    
    try:
        doc = docx.Document(file_path)
        content_items = []
        
        for elem in doc.element.body:
            tag = elem.tag.split('}')[-1]
            
            if tag == 'p':
                p = docx.text.paragraph.Paragraph(elem, doc)
                text = p.text.strip()
                style_name = p.style.name if p.style else 'Normal'
                
                if text:
                    content_items.append({
                        "type": "paragraph",
                        "text": text,
                        "style": style_name
                    })
                
                for img_blob in extract_element_images(elem, doc):
                    content_items.append({
                        "type": "image",
                        "data": img_blob,
                        "format": "png"
                    })
                        
            elif tag == 'tbl':
                table = docx.table.Table(elem, doc)
                table_data = []
                for row in table.rows:
                    row_cells = []
                    for cell in row.cells:
                        for img_blob in extract_element_images(cell._tc, doc):
                            content_items.append({
                                "type": "image",
                                "data": img_blob,
                                "format": "png"
                            })
                        row_cells.append(cell.text.strip().replace('\n', ' '))
                    table_data.append(row_cells)
                if table_data:
                    content_items.append({
                        "type": "table",
                        "data": table_data
                    })
            
            elif tag == 'sdt':
                for sub_p in elem.xpath('.//*[local-name()="p"]'):
                    try:
                        p = docx.text.paragraph.Paragraph(sub_p, doc)
                        text = p.text.strip()
                        if text:
                            content_items.append({
                                "type": "paragraph",
                                "text": text,
                                "style": p.style.name if p.style else 'Normal'
                            })
                    except Exception:
                        pass
                for img_blob in extract_element_images(elem, doc):
                    content_items.append({
                        "type": "image",
                        "data": img_blob,
                        "format": "png"
                    })
                    
        return {"items": content_items, "title": os.path.basename(file_path)}
    except Exception as e:
        return {"error": str(e)}

@st.cache_data(show_spinner=False)
def parse_excel(file_path: str):
    try:
        xl = pd.ExcelFile(file_path)
        sheets_data = {}
        for sheet_name in xl.sheet_names:
            try:
                df = xl.parse(sheet_name)
                # Sanitize column headers and mixed data types to avoid Arrow conversion errors
                df.columns = [str(c).strip() for c in df.columns]
                for col in df.columns:
                    if df[col].dtype == 'object':
                        df[col] = df[col].astype(str).replace('nan', '')
                sheets_data[sheet_name] = df
            except Exception as e:
                sheets_data[sheet_name] = pd.DataFrame({"Error": [f"Could not parse sheet: {e}"]})
        return {"sheets": sheets_data}
    except Exception as e:
        return {"error": str(e)}

@st.cache_data(show_spinner=False)
def parse_pptx(file_path: str):
    if not PPTX_AVAILABLE:
        return {"error": "python-pptx library is not installed."}
    try:
        prs = Presentation(file_path)
        slides = []
        for i, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        slide_texts.append(text)
            slides.append({
                "slide_num": i + 1,
                "text_content": slide_texts
            })
        return {"slides": slides}
    except Exception as e:
        return {"error": str(e)}

# ---------------------------------------------------------
# Environment & Paths
# ---------------------------------------------------------
SAMPLE_DOCS_DIR = os.path.join(os.path.dirname(__file__), "sample_docs")
IS_CLOUD_MODE = not os.path.exists(DEFAULT_DOCS_ROOT)

# ---------------------------------------------------------
# State Management
# ---------------------------------------------------------
if "selected_file_path" not in st.session_state:
    st.session_state["selected_file_path"] = None

def select_document(file_path):
    st.session_state["selected_file_path"] = file_path

def go_home():
    st.session_state["selected_file_path"] = None

def reset_filters_callback():
    st.session_state["filter_search"] = ""
    st.session_state["filter_category"] = "All Categories"
    st.session_state["filter_ext"] = "All Types"

# ---------------------------------------------------------
# Sidebar Controls & Filters
# ---------------------------------------------------------
st.sidebar.markdown("## 📚 Documentation Portal")

# Navigation home button
if st.sidebar.button("🏠 Home / All Categories", width="stretch", type="primary" if st.session_state["selected_file_path"] is None else "secondary"):
    go_home()
    st.rerun()

st.sidebar.markdown("---")
st.sidebar.markdown("### 🔍 Filter Documents")

# Search Input
search_term = st.sidebar.text_input("Search keywords:", placeholder="e.g. Austin, QC, Invoice...", key="filter_search")

# Folder scanning
if os.path.exists(DEFAULT_DOCS_ROOT):
    docs_root = DEFAULT_DOCS_ROOT
    is_cloud_mode = False
else:
    docs_root = SAMPLE_DOCS_DIR
    is_cloud_mode = True

if not os.path.exists(docs_root):
    os.makedirs(docs_root, exist_ok=True)

if is_cloud_mode:
    st.sidebar.caption("☁️ *Browsing cloud sample archives*")

all_files = scan_documents_folder(docs_root)
if not all_files:
    st.warning("No files found.")
    st.stop()

# Category Filter
categories_list = ["All Categories"] + sorted(list({f["category"] for f in all_files}))
selected_cat = st.sidebar.selectbox("Category:", categories_list, key="filter_category")

# Extension Filter
ext_options = {
    "All Types": None,
    "Word Documents (.docx)": [".docx"],
    "Excel Spreadsheets (.xlsx, .xls)": [".xlsx", ".xls"],
    "PowerPoint (.pptx)": [".pptx"]
}
selected_ext_label = st.sidebar.selectbox("File Type:", list(ext_options.keys()), key="filter_ext")
allowed_exts = ext_options[selected_ext_label]

# Apply Filters
filtered_files = all_files

# 1. Filter by category
if selected_cat != "All Categories":
    filtered_files = [f for f in filtered_files if f["category"] == selected_cat]

# 2. Filter by extension
if allowed_exts is not None:
    filtered_files = [f for f in filtered_files if f["ext"] in allowed_exts]

# 3. Filter by search query
if search_term and search_term.strip():
    q = search_term.strip().lower()
    filtered_files = [
        f for f in filtered_files
        if q in f["name"].lower() or q in f["category"].lower() or q in f["rel_path"].lower()
    ]

# Display filter count & reset button
col_cnt, col_rst = st.sidebar.columns([2, 1])
with col_cnt:
    st.caption(f"**{len(filtered_files)}** of {len(all_files)} documents")
with col_rst:
    st.button("Reset", key="reset_filters", on_click=reset_filters_callback)

if st.sidebar.button("🔄 Check for New Files", width="stretch", key="refresh_folder_btn"):
    scan_documents_folder.clear()
    parse_docx.clear()
    parse_excel.clear()
    parse_pptx.clear()
    st.toast("Folder refreshed with latest files!", icon="✅")
    st.rerun()

# ---------------------------------------------------------
# In-App Document Uploader
# ---------------------------------------------------------
with st.sidebar.expander("📤 Upload New Document", expanded=False):
    uploaded_file = st.file_uploader(
        "Select document to add:",
        type=["docx", "xlsx", "xls", "pptx", "pdf", "txt"],
        key="sidebar_doc_uploader"
    )
    existing_cats = sorted(list({f["category"] for f in all_files}))
    upload_cat = st.selectbox("Category:", ["General Procedures"] + existing_cats, key="upload_cat_sel")
    
    if uploaded_file is not None:
        if st.button("💾 Save to Portal", type="primary", key="save_uploaded_doc_btn", width="stretch"):
            target_cat_dir = os.path.join(docs_root, upload_cat)
            os.makedirs(target_cat_dir, exist_ok=True)
            saved_dest = os.path.join(target_cat_dir, uploaded_file.name)
            with open(saved_dest, "wb") as f_save:
                f_save.write(uploaded_file.getbuffer())
            scan_documents_folder.clear()
            st.success(f"✅ Added {uploaded_file.name}!")
            time.sleep(1)
            st.rerun()

st.sidebar.markdown("---")

# ---------------------------------------------------------
# Main Page Area
# ---------------------------------------------------------
current_path = st.session_state["selected_file_path"]
active_doc = next((f for f in all_files if f["full_path"] == current_path), None)

if active_doc is None:
    # -----------------------------------------------------
    # OVERVIEW / DASHBOARD (NOW FULLY REACTIVE TO FILTERS!)
    # -----------------------------------------------------
    share_badge_text = "📁 Active Share: Procedures" if not is_cloud_mode else "☁️ Cloud Mode: Procedures Archive"
    st.markdown(f"""
    <div class='portal-hero'>
        <div style='display:flex; justify-content:space-between; align-items:center;'>
            <div>
                <h1 style='margin:0; font-size:1.85rem; color:#FFFFFF;'>🏢 MLS & Research Operations</h1>
                <p style='margin:6px 0 0 0; color:#DBEAFE; font-size:0.95rem;'>Centralized Standard Operating Procedures (SOPs), Workflows & Quality Checklists</p>
            </div>
            <div style='text-align:right;'>
                <span style='background:rgba(255,255,255,0.2); padding:6px 14px; border-radius:20px; font-size:0.85rem; color:#FFFFFF;'>
                    {share_badge_text}
                </span>
            </div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    # Metrics row (updates with filters)
    m1, m2, m3, m4 = st.columns(4)
    with m1:
        st.markdown(f"<div class='metric-card'><h2 style='margin:0;color:#1E3A8A;'>{len(filtered_files)}</h2><p style='margin:0;color:#64748B;'>Matching Docs</p></div>", unsafe_allow_html=True)
    with m2:
        n_docx = sum(1 for f in filtered_files if f['ext'] == '.docx')
        st.markdown(f"<div class='metric-card'><h2 style='margin:0;color:#2563EB;'>{n_docx}</h2><p style='margin:0;color:#64748B;'>Word Manuals</p></div>", unsafe_allow_html=True)
    with m3:
        n_xl = sum(1 for f in filtered_files if f['ext'] in ['.xlsx', '.xls'])
        st.markdown(f"<div class='metric-card'><h2 style='margin:0;color:#16A34A;'>{n_xl}</h2><p style='margin:0;color:#64748B;'>Excel Checklists</p></div>", unsafe_allow_html=True)
    with m4:
        n_cat = len({f['category'] for f in filtered_files})
        st.markdown(f"<div class='metric-card'><h2 style='margin:0;color:#D97706;'>{n_cat}</h2><p style='margin:0;color:#64748B;'>Categories</p></div>", unsafe_allow_html=True)

    st.markdown("---")

    # If filters applied, show active filter notice
    is_filtered = (selected_cat != "All Categories") or (allowed_exts is not None) or (bool(search_term and search_term.strip()))
    if is_filtered:
        st.info(f"Showing results for: **Category:** {selected_cat} | **Type:** {selected_ext_label} | **Search:** '{search_term or 'None'}'")

    if not filtered_files:
        st.warning("No documents match your filter criteria. Try resetting the filters in the sidebar.")
    else:
        # Group filtered files by category
        filtered_categories = sorted(list({f["category"] for f in filtered_files}))
        for cat in filtered_categories:
            cat_docs = [f for f in filtered_files if f["category"] == cat]
            
            with st.expander(f"📁 **{cat}** — ({len(cat_docs)} documents)", expanded=True if is_filtered else False):
                for doc in cat_docs:
                    c_badge, c_title, c_size, c_btn = st.columns([1, 6, 2, 2])
                    
                    with c_badge:
                        ext_tag = doc['ext'][1:].upper()
                        st.markdown(f"<span class='badge badge-{doc['ext'][1:]}'>{ext_tag}</span>", unsafe_allow_html=True)
                    with c_title:
                        st.write(f"**{doc['name']}**")
                    with c_size:
                        if doc["size_mb"] >= 1.0:
                            st.caption(f"{doc['size_mb']:.1f} MB")
                        else:
                            st.caption(f"{doc['size_kb']:.1f} KB")
                    with c_btn:
                        if st.button("📖 View Doc", key=f"main_view_{doc['full_path']}", width="stretch"):
                            select_document(doc["full_path"])
                            st.rerun()

else:
    # -----------------------------------------------------
    # SINGLE DOCUMENT READER VIEW
    # -----------------------------------------------------
    # Navigation breadcrumbs & Download action
    col_back, col_title, col_dl = st.columns([1.5, 6, 2.5])
    with col_back:
        if st.button("⬅️ Back to All", width="stretch"):
            go_home()
            st.rerun()
    with col_title:
        st.markdown(f"<div class='main-title'>{active_doc['name']}</div>", unsafe_allow_html=True)
    with col_dl:
        try:
            with open(active_doc["full_path"], "rb") as f_dl:
                file_bytes = f_dl.read()
            st.download_button(
                label="⬇️ Download File",
                data=file_bytes,
                file_name=active_doc["name"],
                mime=get_mime_type(active_doc["ext"]),
                width="stretch",
                type="primary",
                key="doc_view_download_btn"
            )
        except Exception as err:
            st.error(f"Cannot read file for download: {err}")

    # Document Metadata Box & Quick Tools
    st.markdown(f"""
    <div class='doc-meta'>
        <strong>📁 Category:</strong> {active_doc['category']} &nbsp;|&nbsp;
        <strong>📂 Folder:</strong> {active_doc['sub_category']} &nbsp;|&nbsp;
        <strong>⚖️ Size:</strong> {active_doc['size_mb']:.2f} MB ({active_doc['size_kb']:.0f} KB) &nbsp;|&nbsp;
        <strong>🕒 Last Modified:</strong> {active_doc['mod_time']}
        <br>
        <span style='font-size:0.85rem; color:#475569;'><strong>📍 Network Path:</strong> {active_doc['full_path']}</span>
    </div>
    """, unsafe_allow_html=True)

    if os.name == 'nt' and not is_cloud_mode:
        col_save, col_exp = st.columns([1, 1])
        with col_save:
            if st.button("💾 Save directly to my Downloads Folder", width="stretch", key="copy_to_downloads_btn"):
                try:
                    user_dl = os.path.join(os.path.expanduser("~"), "Downloads")
                    dest_file = os.path.join(user_dl, active_doc["name"])
                    shutil.copy2(active_doc["full_path"], dest_file)
                    st.success(f"✅ Successfully saved to: {dest_file}")
                except Exception as e:
                    st.error(f"Could not save file: {e}")
        with col_exp:
            if st.button("📂 Locate in Windows File Explorer", width="stretch", key="open_explorer_btn"):
                try:
                    os.system(f'explorer.exe /select,"{active_doc["full_path"]}"')
                    st.toast("Opening file in File Explorer...")
                except Exception as e:
                    st.error(f"Could not open explorer: {e}")
    else:
        if st.button("💾 Save to my Downloads Folder", width="stretch", key="copy_to_downloads_btn"):
            try:
                user_dl = os.path.join(os.path.expanduser("~"), "Downloads")
                dest_file = os.path.join(user_dl, active_doc["name"])
                shutil.copy2(active_doc["full_path"], dest_file)
                st.success(f"✅ Successfully saved to: {dest_file}")
            except Exception as e:
                st.info("Tip: Use the '⬇️ Download File' button above to save this document to your device.")

    ext = active_doc["ext"]

    # 1. Word Document Rendering
    if ext == ".docx":
        with st.spinner("Rendering document..."):
            parsed = parse_docx(active_doc["full_path"])
            
        if "error" in parsed:
            st.error(f"Error loading document: {parsed['error']}")
        else:
            items = parsed.get("items", [])
            if not items:
                st.info("No content elements could be read from this document.")
            
            for item in items:
                item_type = item.get("type")
                if item_type == "paragraph":
                    style = item.get("style", "").lower()
                    text = item.get("text", "")
                    
                    if "heading 1" in style or "title" in style:
                        st.markdown(f"## {text}")
                    elif "heading 2" in style:
                        st.markdown(f"### {text}")
                    elif "heading 3" in style:
                        st.markdown(f"#### {text}")
                    elif "list" in style:
                        st.markdown(f"- {text}")
                    else:
                        st.markdown(text)
                        
                elif item_type == "image":
                    st.image(item["data"], caption="📸 Procedure Screenshot / Diagram", use_container_width=True)
                    
                elif item_type == "table":
                    data = item.get("data", [])
                    if data:
                        max_cols = max(len(r) for r in data)
                        if max_cols > 0:
                            padded_data = [r + [""] * (max_cols - len(r)) for r in data]
                            header = padded_data[0]
                            rows = padded_data[1:] if len(padded_data) > 1 else []
                            
                            clean_header = []
                            for idx, h in enumerate(header):
                                col_name = str(h).strip() if str(h).strip() else f"Col_{idx+1}"
                                while col_name in clean_header:
                                    col_name = f"{col_name}_{idx+1}"
                                clean_header.append(col_name)
                                
                            try:
                                if rows:
                                    df_table = pd.DataFrame(rows, columns=clean_header)
                                else:
                                    df_table = pd.DataFrame([clean_header], columns=clean_header)
                                st.dataframe(df_table, use_container_width=True, hide_index=True)
                            except Exception:
                                st.table(padded_data)

    # 2. Excel Rendering
    elif ext in [".xlsx", ".xls"]:
        with st.spinner("Loading workbook..."):
            parsed = parse_excel(active_doc["full_path"])
            
        if "error" in parsed:
            st.error(f"Error loading workbook: {parsed['error']}")
        else:
            sheets = parsed.get("sheets", {})
            sheet_names = list(sheets.keys())
            
            st.markdown(f"**Workbook Sheets ({len(sheet_names)}):**")
            tabs = st.tabs(sheet_names)
            for i, sheet_name in enumerate(sheet_names):
                with tabs[i]:
                    df_sheet = sheets[sheet_name]
                    c_rows, c_cols = df_sheet.shape
                    
                    c_info, c_view = st.columns([2, 1])
                    with c_info:
                        st.caption(f"Rows: **{c_rows}** | Columns: **{c_cols}**")
                    with c_view:
                        view_full = st.checkbox("Expand Full Table", key=f"chk_full_{active_doc['name']}_{sheet_name}")
                    
                    if view_full:
                        st.dataframe(df_sheet, use_container_width=True, height=min(800, max(250, c_rows * 36)))
                    else:
                        st.dataframe(df_sheet, use_container_width=True)

    # 3. PowerPoint Rendering
    elif ext == ".pptx":
        with st.spinner("Loading slides..."):
            parsed = parse_pptx(active_doc["full_path"])
            
        if "error" in parsed:
            st.error(f"Error loading presentation: {parsed['error']}")
        else:
            slides = parsed.get("slides", [])
            st.markdown(f"**Total Slides: {len(slides)}**")
            for slide in slides:
                with st.expander(f"🎞️ Slide {slide['slide_num']}", expanded=True):
                    content = slide.get("text_content", [])
                    if content:
                        for text in content:
                            if len(text.split('\n')) > 1:
                                for line in text.split('\n'):
                                    st.markdown(f"- {line}")
                            else:
                                st.markdown(f"- **{text}**")
                    else:
                        st.caption("Visual / diagram slide.")
